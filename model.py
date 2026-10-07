# -*- coding: utf-8 -*-
"""AI inference economics: what does it cost to serve an open-weight LLM on rented GPUs, and what margin do
API hosts make at their list prices?

    py model.py      reads data/*.csv (public, cited, retrieved 6 Oct 2026; see data/SOURCES.md)
                     writes outputs/results.json, outputs/AI Inference Economics.xlsx and outputs/*.svg charts

Method
  Serving cost per request  = GPU $/hour / (output tokens/sec per GPU x 3,600 x utilisation) x output tokens
  List price per request    = input tokens x input $/token + output tokens x output $/token
  Implied gross margin      = 1 - serving cost / list price            (compute only; excludes staff, network, idle reserve)
  Break-even utilisation    = the GPU utilisation at which serving cost equals the list price

Each case uses ONE published benchmark (model, GPU, precision, input/output lengths) and the hosts' list prices for
the same model. Throughput is NVIDIA TensorRT-LLM / MLPerf "offline" maximum-load output throughput, i.e. a ceiling;
real services run below it, which the utilisation assumption (base 60%, range 30-90%) is there to capture.
GPU prices are on-demand list prices, grouped into specialist GPU clouds and hyperscalers.

Author: Tolu Oladuji, October 2026. Python standard library + openpyxl.
"""
import csv, json, os, statistics as st

HERE = os.path.dirname(os.path.abspath(__file__))
DATA, OUT = os.path.join(HERE, "data"), os.path.join(HERE, "outputs")
os.makedirs(OUT, exist_ok=True)
BASE_UTIL = 0.60
UTILS = [0.3, 0.4, 0.5, 0.6, 0.7, 0.8, 0.9]


def load(name):
    with open(os.path.join(DATA, name), encoding="utf-8") as f:
        return list(csv.DictReader(f))


def f(x):
    try:
        return float(str(x).replace("$", "").replace(",", ""))
    except (TypeError, ValueError):
        return None


prices, gpus, bench = load("api_prices.csv"), load("gpu_rental.csv"), load("throughput_benchmarks.csv")

# ---------- GPU price groups (on-demand list prices, bare GPU rental; managed-serving rates and PCIe excluded)
GPU_CLOUDS = ("Lambda", "CoreWeave", "RunPod", "Together AI")
HYPERSCALERS = ("AWS", "Google Cloud", "Microsoft Azure")


def gpu_prices(model_key, group):
    out = []
    for g in gpus:
        if g["pricing_type"].strip().lower() != "on-demand" or g["provider"] not in group:
            continue
        name = g["gpu"].lower()
        if model_key not in name or "pcie" in name or "mega" in name:
            continue
        v = f(g["usd_per_gpu_hour"])
        if v:
            out.append((g["provider"], v))
    seen, res = set(), []
    for p, v in out:                    # one quote per provider (AWS lists the same rate for two instance sizes)
        if p not in seen:
            seen.add(p); res.append((p, v))
    return res


small = load("gpu_rental_small.csv") if os.path.exists(os.path.join(DATA, "gpu_rental_small.csv")) else []


def small_prices(model_key):
    """Small GPU clouds and marketplaces: on-demand list prices, one per provider. Marketplaces use their MEDIAN listing
    (never the lowest); 'from' / 'starting at' floor prices and plans that cannot currently be ordered are excluded;
    H100 must be SXM/HGX (the form factor of the benchmarks)."""
    out, seen = [], set()
    for g in small:
        name, notes = g["gpu"], g["notes"].lower()
        if g["pricing_type"].strip().lower() != "on-demand" or model_key not in name.lower():
            continue
        if g["segment"] == "marketplace" and "median" not in name.lower():
            continue
        if "from $" in notes or "'from" in notes or "starting at" in notes or "deploy_ondemand=false" in notes:
            continue
        if model_key == "h100" and not any(t in name.upper() for t in ("SXM", "HGX")):
            continue
        v = f(g["usd_per_gpu_hour"])
        if v and g["provider"] not in seen:
            seen.add(g["provider"]); out.append((g["provider"], v))
    return out


GPU = {}
for key, label in (("h100", "H100"), ("h200", "H200"), ("b200", "B200")):
    gc, hs, sm = gpu_prices(key, GPU_CLOUDS), gpu_prices(key, HYPERSCALERS), small_prices(key)
    GPU[label] = {"gpu_clouds": gc, "hyperscalers": hs, "small_clouds": sm,
                  "small_cloud_median": st.median([v for _, v in sm]) if sm else None,
                  "gpu_cloud_median": st.median([v for _, v in gc]) if gc else None,
                  "hyperscaler_median": st.median([v for _, v in hs]) if hs else None,
                  "cheapest": min([v for _, v in gc + hs + sm]) if gc or hs or sm else None}


def bench_row(src, model, gpu, i=None, o=None):
    for b in bench:
        if src in b["source"] and model in b["model"] and gpu in b["gpu"] and \
           (i is None or b["input_tokens"] == str(i)) and (o is None or b["output_tokens"] == str(o)):
            return b
    raise SystemExit("benchmark not found: %s %s %s" % (src, model, gpu))


def host_rows(*pairs):
    rows = []
    for prov, key in pairs:
        r = next((p for p in prices if p["provider"] == prov and key in p["model"]), None)
        if r and f(r["input_usd_per_1m_tokens"]) is not None:
            rows.append(r)
    return rows


CASES = [
    {"name": "Llama 3.1 8B", "gpu": "H100", "lengths": (1000, 1000),
     "bench": bench_row("TensorRT-LLM", "Llama 3.1 8B", "H100", 1000, 1000),
     "hosts": host_rows(("DeepInfra", "Llama-3.1-8B"), ("Together AI", "Llama 3 8B"), ("Fireworks AI", "4B-16B"), ("Amazon Bedrock", "8B"))},
    {"name": "Llama 3.3 70B", "gpu": "H100", "lengths": (1000, 1000),
     "bench": bench_row("TensorRT-LLM", "Llama 3.3 70B", "H100", 1000, 1000),
     "hosts": host_rows(("DeepInfra", "Llama-3.3-70B"), ("Amazon Bedrock", "70B"), ("Fireworks AI", ">16B"), ("Together AI", "Llama 3.3 70B"))},
    {"name": "Llama 3.3 70B on next-gen B200", "gpu": "B200", "lengths": (1000, 1000),
     "bench": bench_row("TensorRT-LLM", "Llama 3.3 70B", "B200", 1000, 1000),
     "hosts": host_rows(("DeepInfra", "Llama-3.3-70B"), ("Amazon Bedrock", "70B"), ("Fireworks AI", ">16B"), ("Together AI", "Llama 3.3 70B"))},
    {"name": "gpt-oss-120b (MoE)", "gpu": "B200", "lengths": (None, None),
     "bench": bench_row("MLPerf Inference v6.1", "gpt-oss-120b", "B200"),
     "hosts": host_rows(("Together AI", "gpt-oss-120B"), ("Fireworks AI", "GPT OSS 120B"), ("Groq", "gpt-oss-120b"), ("Amazon Bedrock", "gpt-oss-120b"))},
]

results = []
for c in CASES:
    b = c["bench"]
    tps = f(b["throughput_tokens_per_sec_per_gpu"])
    g = GPU[c["gpu"]]
    i_len, o_len = c["lengths"]
    row = {"case": c["name"], "gpu": c["gpu"], "benchmark": "%s | %s | %s x%s %s | in %s / out %s | %s tok/s per GPU" % (
               b["source"], b["model"], b["gpu"], b["num_gpus"], b["precision"], b["input_tokens"] or "n/a", b["output_tokens"] or "n/a", tps),
           "benchmark_url": b["source_url"], "tps_per_gpu": tps, "gpu_prices": g, "tiers": {}, "hosts": [], "sensitivity": []}
    tiers = {"Cheapest on-demand": g["cheapest"], "Small-cloud / marketplace median": g["small_cloud_median"],
             "GPU-cloud median": g["gpu_cloud_median"], "Hyperscaler median": g["hyperscaler_median"]}
    for tier, gph in tiers.items():
        if gph:
            row["tiers"][tier] = {"usd_per_gpu_hour": gph, "cost_per_1m_output": gph / (tps * 3600 * BASE_UTIL) * 1e6}
    for u in UTILS:
        for tier, gph in tiers.items():
            if gph:
                row["sensitivity"].append({"tier": tier, "util": u, "cost_per_1m_output": gph / (tps * 3600 * u) * 1e6})
    base_cost_m = row["tiers"]["GPU-cloud median"]["cost_per_1m_output"]
    for h in c["hosts"]:
        pi, po = f(h["input_usd_per_1m_tokens"]), f(h["output_usd_per_1m_tokens"])
        if i_len:   # full request: price and cost at the benchmark's own input/output lengths
            price = (i_len * pi + o_len * po) / 1e6
            cost = base_cost_m * o_len / 1e6
            basis = "per request, %d in / %d out" % (i_len, o_len)
        else:       # benchmark lengths not stated: compare cost per 1M output tokens with the OUTPUT price only (conservative)
            price, cost, basis = po, base_cost_m, "per 1M output tokens vs output price only (input revenue ignored)"
        row["hosts"].append({"host": h["provider"], "host_model": h["model"], "input_price": pi, "output_price": po,
                             "price": price, "cost": cost, "basis": basis, "gross_margin": 1 - cost / price,
                             "breakeven_util": BASE_UTIL * cost / price, "source_url": h["source_url"]})
    results.append(row)

closed = [{"provider": p["provider"], "model": p["model"], "input": f(p["input_usd_per_1m_tokens"]), "output": f(p["output_usd_per_1m_tokens"])}
          for p in prices if p["open_weights"].strip().lower() == "no"]

# ---------- the whole market, big and small: every provider offer on OpenRouter for the same three models
# Each offer is costed on the CHEAPEST realistic rented capacity a small host could use (small-cloud / marketplace
# median), so the margins shown are the most favourable a GPU-renting host could earn at that price.
LT = []
if os.path.exists(os.path.join(DATA, "openrouter_prices.csv")):
    bench_by_model = {"Llama 3.1 8B": next(r for r in results if r["case"] == "Llama 3.1 8B"),
                      "Llama 3.3 70B": next(r for r in results if r["case"] == "Llama 3.3 70B"),
                      "gpt-oss-120b": next(r for r in results if r["case"] == "gpt-oss-120b (MoE)")}
    for o in load("openrouter_prices.csv"):
        r = bench_by_model.get(o["model"])
        if not r:
            continue
        tier = r["tiers"].get("Small-cloud / marketplace median") or r["tiers"]["GPU-cloud median"]
        cm = tier["cost_per_1m_output"]
        pi, po = f(o["input_usd_per_1m_tokens"]), f(o["output_usd_per_1m_tokens"])
        if o["model"] == "gpt-oss-120b":
            price, cost = po, cm
        else:
            price, cost = (1000 * pi + 1000 * po) / 1e6, cm * 1000 / 1e6
        if not price:
            continue
        LT.append({"model": o["model"], "provider": o["provider"], "segment": o["segment"], "quantization": o["quantization"],
                   "input_price": pi, "output_price": po, "price": price, "cost": cost,
                   "gross_margin": 1 - cost / price, "breakeven_util": BASE_UTIL * cost / price})
seg_summary = []
for model in ("Llama 3.1 8B", "Llama 3.3 70B", "gpt-oss-120b"):
    for seg in ("hyperscaler / big cloud", "custom chips", "major independent host", "long-tail independent host"):
        xs = [x for x in LT if x["model"] == model and x["segment"] == seg]
        if xs:
            seg_summary.append({"model": model, "segment": seg, "offers": len(xs),
                                "median_output_price": st.median([x["output_price"] for x in xs]),
                                "min_output_price": min(x["output_price"] for x in xs),
                                "max_output_price": max(x["output_price"] for x in xs),
                                "median_gross_margin": st.median([x["gross_margin"] for x in xs]),
                                "offers_below_cost": sum(1 for x in xs if x["gross_margin"] < 0)})

json.dump({"base_utilisation": BASE_UTIL, "gpu": GPU, "cases": results, "closed_model_prices": closed,
           "market_offers": LT, "segment_summary": seg_summary},
          open(os.path.join(OUT, "results.json"), "w", encoding="utf-8"), indent=1)


# ---------- charts (plain SVG, no dependencies)
def svg_bars(path, title, labels, values, fmt, ref=None, ref_label="", colors=None, width=760):
    h_bar, gap, left, top = 26, 12, 290, 50
    vmax = max([abs(v) for v in values] + ([ref] if ref else [])) * 1.15 or 1
    vmin = min(0, min(values))
    span = vmax - vmin
    x0 = left + (0 - vmin) / span * (width - left - 90)
    height = top + len(values) * (h_bar + gap) + 40
    s = ['<svg xmlns="http://www.w3.org/2000/svg" width="%d" height="%d" font-family="Segoe UI,Helvetica,Arial,sans-serif" font-size="13">' % (width, height),
         '<rect width="100%" height="100%" fill="#ffffff"/>',
         '<text x="16" y="28" font-size="16" font-weight="600" fill="#1f2a44">%s</text>' % title]
    for k, (lab, v) in enumerate(zip(labels, values)):
        y = top + k * (h_bar + gap)
        x1 = left + (v - vmin) / span * (width - left - 90)
        col = (colors[k] if colors else ("#2f6fb1" if v >= 0 else "#c0504d"))
        s.append('<text x="%d" y="%d" text-anchor="end" fill="#333">%s</text>' % (left - 10, y + 18, lab))
        s.append('<rect x="%.1f" y="%d" width="%.1f" height="%d" fill="%s" rx="3"/>' % (min(x0, x1), y, abs(x1 - x0), h_bar, col))
        s.append('<text x="%.1f" y="%d" fill="#333">%s</text>' % (max(x0, x1) + 6, y + 18, fmt(v)))
    s.append('<line x1="%.1f" y1="%d" x2="%.1f" y2="%d" stroke="#888"/>' % (x0, top - 6, x0, height - 34))
    if ref:
        xr = left + (ref - vmin) / span * (width - left - 90)
        s.append('<line x1="%.1f" y1="%d" x2="%.1f" y2="%d" stroke="#e08a00" stroke-dasharray="5,4"/>' % (xr, top - 6, xr, height - 34))
        s.append('<text x="%.1f" y="%d" fill="#e08a00" font-size="12">%s</text>' % (xr + 4, height - 18, ref_label))
    s.append("</svg>")
    open(path, "w", encoding="utf-8").write("\n".join(s))


labels, vals = [], []
for r in results[:2]:
    for h in r["hosts"]:
        labels.append("%s | %s" % (r["case"], h["host"].replace("Amazon ", "")))
        vals.append(h["gross_margin"] * 100)
svg_bars(os.path.join(OUT, "gross_margin_by_host.svg"),
         "Implied gross margin at list price (H100, GPU-cloud median rate, 60% utilisation)",
         labels, vals, lambda v: "%.0f%%" % v)
r70 = results[1]
labels = list(r70["tiers"].keys()) + ["B200, GPU-cloud median"]
vals = [t["cost_per_1m_output"] for t in r70["tiers"].values()] + [results[2]["tiers"]["GPU-cloud median"]["cost_per_1m_output"]]
svg_bars(os.path.join(OUT, "cost_per_1m_tokens_70b.svg"),
         "Llama 3.3 70B: GPU cost per 1M output tokens at 60% utilisation", labels, vals, lambda v: "$%.2f" % v,
         ref=1.04, ref_label="Together list price $1.04", colors=["#9dbbe0", "#5b8fc9", "#2f6fb1", "#1f3864", "#3a9d5d"])

# ---------- Excel
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
NAVY = "1F3864"; HF = Font(bold=True, color="FFFFFF"); FILL = PatternFill("solid", fgColor=NAVY)


def table(ws, r0, headers, rows, fmts):
    for j, hd in enumerate(headers, 1):
        c = ws.cell(row=r0, column=j, value=hd); c.font = HF; c.fill = FILL; c.alignment = Alignment(wrap_text=True)
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=i, column=j, value=v)
            if fmts[j - 1]:
                c.number_format = fmts[j - 1]
    return r0 + len(rows) + 2


wb = Workbook(); ws = wb.active; ws.title = "README"; ws.column_dimensions["A"].width = 125
for i, line in enumerate(__doc__.strip().splitlines(), 1):
    ws.cell(row=i, column=1, value=line)
ws = wb.create_sheet("Results")
for col, w in zip("ABCDEFGHI", [34, 34, 12, 12, 16, 16, 14, 16, 40]):
    ws.column_dimensions[col].width = w
r = 1
for res in results:
    ws.cell(row=r, column=1, value=res["case"]).font = Font(bold=True, size=12, color=NAVY)
    ws.cell(row=r + 1, column=1, value=res["benchmark"])
    r = table(ws, r + 2, ["GPU price tier", "", "$/GPU-hour", "", "Cost per 1M output tokens (60% util)"],
              [[t, "", v["usd_per_gpu_hour"], "", v["cost_per_1m_output"]] for t, v in res["tiers"].items()],
              [None, None, "$#,##0.00", None, "$#,##0.000"])
    r = table(ws, r, ["Host", "Host model", "$/1M in", "$/1M out", "Price", "Serving cost", "Gross margin", "Break-even util", "Basis"],
              [[h["host"], h["host_model"], h["input_price"], h["output_price"], h["price"], h["cost"], h["gross_margin"], h["breakeven_util"], h["basis"]] for h in res["hosts"]],
              [None, None, "$#,##0.000", "$#,##0.000", "$#,##0.000000", "$#,##0.000000", "0.0%", "0.0%", None])
ws = wb.create_sheet("Sensitivity")
rows = [[res["case"], s["tier"], s["util"], s["cost_per_1m_output"]] for res in results for s in res["sensitivity"]]
table(ws, 1, ["Case", "GPU price tier", "Utilisation", "Cost per 1M output tokens"], rows, [None, None, "0%", "$#,##0.000"])
for col, w in zip("ABCD", [34, 22, 12, 20]):
    ws.column_dimensions[col].width = w
ws = wb.create_sheet("Whole market (OpenRouter)")
table(ws, 1, ["Model", "Provider", "Segment", "Precision", "$/1M in", "$/1M out", "Price", "Cost (small-cloud GPUs)", "Gross margin", "Break-even util"],
      [[x["model"], x["provider"], x["segment"], x["quantization"], x["input_price"], x["output_price"], x["price"], x["cost"], x["gross_margin"], x["breakeven_util"]] for x in LT],
      [None, None, None, None, "$#,##0.000", "$#,##0.000", "$#,##0.000000", "$#,##0.000000", "0.0%", "0.0%"])
for col, w in zip("ABCDEFGHIJ", [16, 16, 26, 10, 10, 10, 13, 16, 12, 13]):
    ws.column_dimensions[col].width = w
ws = wb.create_sheet("Segment summary")
table(ws, 1, ["Model", "Segment", "Offers", "Median $/1M out", "Min $/1M out", "Max $/1M out", "Median gross margin", "Offers below cost"],
      [[s["model"], s["segment"], s["offers"], s["median_output_price"], s["min_output_price"], s["max_output_price"], s["median_gross_margin"], s["offers_below_cost"]] for s in seg_summary],
      [None, None, "0", "$#,##0.000", "$#,##0.000", "$#,##0.000", "0.0%", "0"])
for col, w in zip("ABCDEFGH", [16, 26, 8, 14, 12, 12, 16, 14]):
    ws.column_dimensions[col].width = w
for name, rows_ in (("GPU prices", gpus), ("API prices", prices), ("Benchmarks", bench)):
    ws = wb.create_sheet(name)
    keys = list(rows_[0].keys())
    table(ws, 1, keys, [[r_.get(k) for k in keys] for r_ in rows_], [None] * len(keys))
wb.save(os.path.join(OUT, "AI Inference Economics.xlsx"))

for res in results:
    print("\n" + res["case"], "|", res["benchmark"])
    for t, v in res["tiers"].items():
        print("   %-20s $%.2f/GPU-hr -> $%.3f per 1M output tokens" % (t, v["usd_per_gpu_hour"], v["cost_per_1m_output"]))
    for h in res["hosts"]:
        print("   %-16s GM %6.1f%%  break-even util %6.1f%%  (%s)" % (h["host"], 100 * h["gross_margin"], 100 * h["breakeven_util"], h["basis"]))
print("\nWHOLE MARKET by segment (costed on small-cloud / marketplace GPUs, %d%% utilisation)" % int(BASE_UTIL * 100))
for s in seg_summary:
    print("   %-14s %-27s n=%-2d out $%.2f (min %.2f, max %.2f)  median GM %6.1f%%  below cost %d" % (
        s["model"], s["segment"], s["offers"], s["median_output_price"], s["min_output_price"], s["max_output_price"], 100 * s["median_gross_margin"], s["offers_below_cost"]))
print("GPU medians:", {k: (round(v["small_cloud_median"] or 0, 2), round(v["gpu_cloud_median"] or 0, 2), round(v["hyperscaler_median"] or 0, 2)) for k, v in GPU.items()})
